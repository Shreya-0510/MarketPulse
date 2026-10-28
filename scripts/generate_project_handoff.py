from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "MarketPulse Project Handoff - Updated.docx"
SCREENSHOTS = ROOT / "docs" / "screenshots"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255))
        shade(t.rows[0].cells[i], "1F4E79")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(t.rows) % 2 == 0:
                shade(cells[i], "EAF2F8")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return t


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.style = "No Spacing"
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    return p


def add_screenshot(doc, filename, title, explanation):
    doc.add_page_break()
    heading = doc.add_heading(title, level=2)
    heading.paragraph_format.keep_with_next = True
    image = SCREENSHOTS / filename
    if image.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        # The documentation assets are pre-cropped to a page-safe viewport.
        # Keep the image below the heading and caption on the same page.
        p.add_run().add_picture(str(image), width=Inches(6.45))
        p.paragraph_format.keep_with_next = True
    caption = doc.add_paragraph(explanation)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("MarketPulse").bold = True
    p.runs[0].font.size = Pt(30)
    p.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Handoff and Demo Guide")
    r.bold = True
    r.font.size = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Real-Time Market Intelligence and Sentiment Analytics").italic = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Prepared for a four-person technical demonstration | October 2026")
    table(doc, ["Document purpose", "Audience", "Runtime status"], [
        ["Explain the system, implementation, file ownership, and demo flow.", "Team members, evaluators, and future maintainers.", "Docker stack verified locally; dashboard and data pipeline operational."],
    ])
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "MarketPulse is a containerized data platform that collects daily market bars, public news, and Bluesky posts; "
        "normalizes and transports them through Kafka; processes them with Spark Structured Streaming; stores them in MongoDB; "
        "and exposes them through FastAPI to a React dashboard. The system emphasizes observable data quality: deterministic event IDs, "
        "idempotent upserts, explicit source status, freshness indicators, and honest empty or unknown states."
    )
    doc.add_paragraph(
        "This document is intentionally organized for a four-person demo. Each person has a coherent ownership area, a file list, "
        "implementation details, a short explanation to memorize, and likely evaluator questions."
    )

    doc.add_heading("1. Problem statement and objectives", level=1)
    doc.add_paragraph(
        "Market information is distributed across price providers, news publishers, and social platforms. A user needs one place "
        "to inspect current watchlist movement, relevant public discussion, sentiment, and pipeline health without confusing missing data "
        "with zero activity."
    )
    for item in [
        "Collect heterogeneous external data using purpose-built adapters.",
        "Use Kafka to decouple producers from downstream processing.",
        "Use Spark for streaming validation, transformation, sentiment, and persistence.",
        "Persist durable records and analytics in MongoDB.",
        "Provide a read-only FastAPI API and a browser dashboard.",
        "Make delivery behavior, freshness, failures, and limitations visible.",
    ]:
        bullet(doc, item)
    doc.add_heading("Non-goals", level=2)
    for item in ["Trading or order execution", "Investment advice", "Price prediction", "Scraping full article bodies", "Exactly-once delivery guarantees"]:
        bullet(doc, item)

    doc.add_heading("2. End-to-end architecture", level=1)
    code(doc, "yfinance -> ingestion-market --\\")
    code(doc, "Bluesky -> ingestion-social ----+-> Kafka -> Spark -> MongoDB -> FastAPI -> nginx -> React")
    code(doc, "RSS/Google News -> ingestion-news -/                 ^")
    code(doc, "                                      analytics-job --/")
    doc.add_paragraph(
        "The adapters publish events and do not write directly to application collections. Kafka provides buffering and decoupling. "
        "Spark consumers validate and persist events. The API reads persisted data only, which keeps dashboard requests deterministic "
        "and prevents external provider latency from leaking into the user experience."
    )
    table(doc, ["Layer", "Technology", "Responsibility"], [
        ["Sources", "yfinance, RSS/Atom, Google News RSS, Bluesky Jetstream", "Fetch or receive external records."],
        ["Transport", "Apache Kafka 3.9 KRaft", "Topics, buffering, and producer/consumer decoupling."],
        ["Processing", "PySpark 3.5.3", "Validation, deduplication, sentiment, and Mongo sinks."],
        ["Storage", "MongoDB 7", "Durable event collections and analytics collections."],
        ["API", "FastAPI + Pydantic", "Typed read APIs, health, source status, and WebSocket feed."],
        ["UI", "React + TypeScript + Vite + Tailwind + Recharts", "Four dashboard workspaces and explicit data states."],
        ["Runtime", "Docker Compose + nginx", "Repeatable local orchestration and same-origin proxying."],
    ])

    doc.add_heading("3. Repository and file structure", level=1)
    code(doc, """MarketPulse/
  apps/web/                 React dashboard
  services/common/          shared contracts, IDs, config, matching, Mongo schema
  services/ingestion/       market, RSS/news, Bluesky adapters
  services/streaming/       Spark jobs, transforms, sinks, sentiment, analytics
  services/api/              FastAPI routes, repositories, schemas
  config/                    watchlist, entities, sources, settings
  infra/compose.yaml         service graph
  infra/docker/              image definitions
  scripts/                   startup, smoke tests, validation, verification
  docs/screenshots/          dashboard evidence
  BUILD.md                   detailed operational contracts
  README.md                  public project documentation""")
    table(doc, ["Path", "Important contents", "Why it matters"], [
        ["services/common/marketpulse_common/events.py", "Pydantic event models", "Single contract shared across producers and consumers."],
        ["services/common/marketpulse_common/ids.py", "Stable event ID helpers", "Makes retries and redelivery idempotent."],
        ["services/ingestion/mp_ingestion/market.py", "yfinance fetch and normalization", "Preserves missingness and rejects invalid bars."],
        ["services/ingestion/mp_ingestion/news.py", "Feed parsing and conditional GET", "Normalizes RSS/Atom without article scraping."],
        ["services/ingestion/mp_ingestion/social.py", "Jetstream parser", "Keeps matching Bluesky post creates only."],
        ["services/streaming/mp_streaming/sinks.py", "Market Mongo sink", "Upserts market events and records Kafka trace metadata."],
        ["services/streaming/mp_streaming/sinks_text.py", "News/social sinks", "Persists text events with delivery metadata."],
        ["services/streaming/mp_streaming/analytics_job.py", "Scheduled analytics", "Scores text and computes entity/attention/anomaly metrics."],
        ["services/api/app/routes/", "Market, text, analytics, operations routes", "Stable read surface for the dashboard."],
        ["apps/web/src/pages/", "Overview, EntityDeepDive, Firehose, Pipeline", "User-facing demo workspaces."],
        ["infra/compose.yaml", "All containers, dependencies, health checks", "Defines the reproducible local environment."],
    ])

    doc.add_heading("4. Implementation details", level=1)
    doc.add_heading("Event contracts and identity", level=2)
    doc.add_paragraph(
        "Pydantic models define MarketBarEvent, NewsArticleEvent, and SocialPostEvent. Each event carries source, event time, ingestion time, "
        "data mode, entity IDs, and source metadata. IDs are derived from stable provider identity instead of random UUIDs. For example, a market "
        "ID is based on provider, symbol, interval, and timestamp; a news ID uses canonical URL/publisher/title; a social ID uses the Bluesky URI."
    )
    doc.add_heading("Market ingestion", level=2)
    numbered(doc, "Load the watchlist and source settings.")
    numbered(doc, "Use a one-year initial lookback, then an incremental period when Mongo already has history.")
    numbered(doc, "Fetch OHLCV rows through yfinance with bounded retries.")
    numbered(doc, "Reject naive timestamps, invalid closes, NaN, infinity, and contract violations.")
    numbered(doc, "Publish normalized records to market.raw and record per-symbol status.")
    doc.add_heading("News ingestion", level=2)
    doc.add_paragraph(
        "The RSS adapter uses conditional GET headers (ETag and Last-Modified), follows redirects, retries failures, and parses feeds with feedparser. "
        "It keeps headlines, summaries, links, publisher metadata, publication time, and entity matches. The retired Yahoo Finance RSS endpoint is disabled; "
        "Google News entity RSS search is enabled as the supported replacement. A malformed feed can remain partially useful and is reported in source status."
    )
    doc.add_heading("Social ingestion", level=2)
    doc.add_paragraph(
        "The Bluesky Jetstream adapter receives commit messages over WebSocket, retains only create operations for app.bsky.feed.post, matches configured "
        "entities, clamps untrusted future createdAt values, and publishes matching posts. The public stream is intentionally filtered because most posts "
        "are irrelevant to the watchlist."
    )
    doc.add_heading("Spark processing and MongoDB", level=2)
    doc.add_paragraph(
        "Spark Structured Streaming reads Kafka topics with checkpoints under /checkpoints. Validation and transforms produce typed rows. Mongo sinks use "
        "upserts keyed by event_id, so at-least-once delivery is safe for the persisted view. News and social text is scored with VADER by analytics-job. "
        "Production analytics excludes data_mode=test records used by smoke tests."
    )
    doc.add_heading("API and frontend", level=2)
    doc.add_paragraph(
        "FastAPI repositories are read-only and convert Mongo documents into typed response schemas. The web image builds the Vite SPA and nginx proxies "
        "/api and WebSocket traffic to the API container. The frontend polls health, sources, watchlist, news, social, attention, anomalies, and entity metrics; "
        "the Firehose maintains one WebSocket for live persisted-event updates."
    )

    doc.add_heading("5. Team ownership split", level=1)
    table(doc, ["Person", "Ownership", "Primary demo story"], [
        ["Person 1", "Data sources, configuration, and ingestion", "How external data becomes validated Kafka events."],
        ["Person 2", "Kafka, Spark streaming, MongoDB, and reliability", "How events are processed, deduplicated, and persisted."],
        ["Person 3", "FastAPI backend and analytics", "How persisted data becomes typed REST/WebSocket responses and metrics."],
        ["Person 4", "React dashboard, Docker orchestration, and QA", "How the user sees the data and how the stack is run and verified."],
    ])

    assignments = [
        (
            "Person 1 — Sources and ingestion",
            ["config/sources.yaml", "config/entities.yaml", "config/watchlist.yaml", "services/ingestion/mp_ingestion/market.py",
             "services/ingestion/mp_ingestion/news.py", "services/ingestion/mp_ingestion/social.py",
             "services/ingestion/mp_ingestion/run_market.py", "services/ingestion/mp_ingestion/run_news.py",
             "services/ingestion/mp_ingestion/run_social.py", "services/ingestion/mp_ingestion/kafka_out.py"],
            "You own the boundary between the outside world and MarketPulse. The source configuration in "
            "config/sources.yaml declares the adapter type, URL, polling interval, timeout, and provider-specific options. "
            "config/watchlist.yaml defines the instruments, while config/entities.yaml defines the terms used to decide whether a "
            "news headline or social post is relevant. This separation means a source can be changed without rewriting the processing jobs.\n\n"
            "The market adapter calls yfinance one symbol at a time. It chooses a one-year lookback for a new symbol and a shorter "
            "incremental period after history exists. normalize_rows preserves missing values, rejects invalid timestamps and closes, "
            "and creates a deterministic event_id from provider, symbol, interval, and timestamp. run_market retries a failed symbol "
            "without stopping the remaining watchlist and records per-symbol status in MongoDB.\n\n"
            "The news adapter uses conditional HTTP requests with ETag/Last-Modified validators, follows redirects, retries bounded "
            "failures, parses RSS/Atom entries, strips HTML, keeps source publication time, and matches entities without scraping article "
            "bodies. The old Yahoo Finance RSS URL returned 404/429, so it is disabled and the supported Google News entity adapter is enabled. "
            "The social adapter listens to Bluesky Jetstream, retains post-create events, matches configured terms, validates timestamps, and "
            "publishes only relevant posts. All three adapters publish to Kafka rather than writing application collections directly.\n\n"
            "Demo script: 'I will show how sources.yaml controls adapters, then explain that each adapter normalizes different provider formats "
            "into the same typed event contract. A provider failure is isolated and visible in /api/v1/sources/status; it does not become fake "
            "data or stop the other sources.'",
            ["Why does the dashboard not call providers directly?", "How are duplicate news or market records prevented?", "What happens when one symbol or feed fails?", "Why is Bluesky not representative of all investors?"],
        ),
        (
            "Person 2 — Kafka, Spark, MongoDB, and reliability",
            ["infra/compose.yaml", "infra/docker/streaming.Dockerfile", "services/streaming/mp_streaming/job_runner.py",
             "services/streaming/mp_streaming/market_job.py", "services/streaming/mp_streaming/text_jobs.py",
             "services/streaming/mp_streaming/transforms.py", "services/streaming/mp_streaming/transforms_text.py",
             "services/streaming/mp_streaming/sinks.py", "services/streaming/mp_streaming/sinks_text.py",
             "services/common/marketpulse_common/mongo_schema.py"],
            "You own the durable data path after ingestion. Kafka runs in KRaft mode and the Compose file creates the raw topics "
            "market.raw, social.raw, and news.raw plus analytics topics. kafka-init creates them idempotently after the broker health check. "
            "The ingestion containers publish JSON event payloads, so producers and consumers can restart independently.\n\n"
            "spark-streaming reads market.raw and spark-text reads social.raw/news.raw through Structured Streaming. The job runner applies "
            "a schema, validates records, transforms them into sink rows, and uses checkpoints under /checkpoints so a restart can resume from "
            "known Kafka progress. The Mongo sinks upsert by event_id and attach topic, partition, offset, delivery count, and processing time. "
            "This is at-least-once delivery: a message can be delivered again, but deterministic IDs and upserts make the persisted view safe "
            "from duplicate business records. It is not exactly-once processing.\n\n"
            "analytics-job is separate from the Kafka streaming consumers. It reads persisted market, news, and social collections, scores "
            "unscored text with VADER, writes entity_metrics and attention_metrics, and records anomaly events only when enough history exists. "
            "Mongo schema initialization creates indexes and retention-related structures used by the runtime.\n\n"
            "Demo script: 'I will trace one record from a raw Kafka topic through a Spark micro-batch into MongoDB. The important reliability "
            "choice is not claiming exactly once; instead, the system is honest about at-least-once delivery and makes redelivery harmless with "
            "stable event IDs and idempotent upserts.'",
            ["Why use Kafka between adapters and Spark?", "What does a checkpoint protect against?", "Is the system exactly once?", "How can a persisted event be traced back to Kafka?"],
        ),
        (
            "Person 3 — API and analytics",
            ["services/api/app/main.py", "services/api/app/deps.py", "services/api/app/schemas.py",
             "services/api/app/routes/market.py", "services/api/app/routes/text.py", "services/api/app/routes/analytics.py",
             "services/api/app/routes/ops.py", "services/api/app/repositories/", "services/streaming/mp_streaming/analytics_job.py",
             "services/streaming/mp_streaming/sentiment.py"],
            "You own the read model presented to the browser. app/main.py creates the FastAPI application and exception handling. deps.py "
            "creates the Mongo client and repository dependencies. The route modules separate responsibilities: market.py serves watchlists, "
            "history, and event traces; text.py serves news and social records with filters; analytics.py serves entity metrics, attention, "
            "anomalies, and cross-signals; ops.py serves pipeline and source status.\n\n"
            "Pydantic response schemas in schemas.py prevent the API from leaking arbitrary Mongo documents. Repository classes contain the "
            "database queries, so route functions focus on HTTP validation and response construction. The API reads persisted data only; "
            "provider calls stay inside ingestion services. /health proves that the process is running and Mongo responds. It deliberately "
            "reports Kafka and Spark as UNKNOWN because those checks are not instrumented there. /pipeline/status and /sources/status expose "
            "recorded component documents written by adapters and jobs. Disabled sources are reported as DISABLED even if an old failure "
            "document remains in Mongo.\n\n"
            "analytics_job.py computes closed windows for each configured entity. It counts social and news mentions, separates scored from "
            "unscored text, computes average VADER scores, and compares attention against a baseline for anomaly detection. Cross-signals are "
            "descriptive combinations such as news activity with stable price; they do not imply causation or predict returns. Test records are "
            "excluded from production queries through data_mode filtering.\n\n"
            "Demo script: 'I will open the API endpoints beside the dashboard. The frontend receives typed persisted data, while source health "
            "and freshness explain whether a missing value means no data, stale data, a disabled adapter, or an uninstrumented subsystem.'",
            ["Why does /health report Kafka and Spark as UNKNOWN?", "Why can an analytics chart be empty while ingestion works?", "How are test records excluded?", "What does the cross-signal output claim—and not claim?"],
        ),
        (
            "Person 4 — Frontend, orchestration, and verification",
            ["apps/web/src/App.tsx", "apps/web/src/api/client.ts", "apps/web/src/pages/Overview.tsx",
             "apps/web/src/pages/EntityDeepDive.tsx", "apps/web/src/pages/Firehose.tsx", "apps/web/src/pages/Pipeline.tsx",
             "apps/web/src/hooks/", "apps/web/src/components/", "infra/docker/web.Dockerfile", "scripts/up.ps1",
             "scripts/verify_stack.py", "scripts/smoke_market.py", "scripts/smoke_text.py"],
            "You own the user experience and the repeatable demonstration. App.tsx defines the four hash-routed workspaces and keeps "
            "one live-feed connection for the application lifetime. api/client.ts centralizes typed GET calls, query-string construction, "
            "HTTP error conversion, and the WebSocket URL. The hooks load health, entities, source status, metrics, and live events without "
            "putting provider logic in the browser.\n\n"
            "Overview.tsx composes watchlist, indexed price performance, activity, sentiment, attention, headlines, posts, and source health. "
            "EntityDeepDive.tsx loads a selected entity and its price and analytics windows. Firehose.tsx displays the WebSocket state, server "
            "heartbeat, replayed snapshots, newly arrived events, and filters. Pipeline.tsx explains recorded status and makes UNKNOWN explicit. "
            "Loading, error, empty, stale, and disabled states are intentional; the UI never replaces missing data with zero.\n\n"
            "The web Dockerfile builds the Vite bundle and serves it with nginx. nginx serves the SPA and proxies /api and WebSocket traffic to "
            "api:8000, so the browser can use same-origin URLs. scripts/up.ps1 creates .env and starts Compose. verify_stack.py checks API, source "
            "status, analytics, WebSocket, and web proxy; smoke_market.py and smoke_text.py test real paths and labelled test data.\n\n"
            "Demo script: 'I will start with docker compose, show container health, then use the four workspaces to prove the same persisted "
            "records are visible through REST, WebSocket, and the browser. I will also show that a deliberate empty state is a truthful result, "
            "not a frontend failure.'",
            ["How does the browser reach the API?", "Why are empty analytics panels not errors?", "What does the Firehose snapshot represent?", "How do you prove the complete path works?"],
        ),
    ]
    for title, files, story, questions in assignments:
        doc.add_heading(title, level=2)
        doc.add_paragraph("Files to study:")
        for f in files:
            bullet(doc, f, level=1)
        doc.add_paragraph("Demo explanation: " + story)
        doc.add_paragraph("Likely questions:")
        for q in questions:
            bullet(doc, q, level=1)

    doc.add_heading("6. Demo runbook", level=1)
    for step in [
        "Start Docker Desktop and run Copy-Item .env.example .env followed by docker compose -f infra/compose.yaml --env-file .env up -d --build.",
        "Run docker compose -f infra/compose.yaml ps and show Kafka/Mongo health plus API and web containers.",
        "Open http://localhost:3000 and point out API ok, Mongo ok, and Data FRESH.",
        "Use Overview to explain the watchlist and persisted daily-bar semantics.",
        "Open Entity Deep Dive, select Apple, and explain price context and entity-level metrics.",
        "Open Live Firehose and show WebSocket open, heartbeat, filters, and replay/new-event semantics.",
        "Open Pipeline & Methodology and explain why non-instrumented Kafka/Spark fields say UNKNOWN.",
        "Show /api/v1/sources/status and explain source failures are isolated and visible.",
        "If asked about analytics empty states, explain the five-minute analytics cycle and anomaly history requirement.",
    ]:
        numbered(doc, step)

    doc.add_heading("7. Dashboard screenshots", level=1)
    add_screenshot(doc, "doc-overview.png", "Overview — watchlist and platform health",
                   "The Overview workspace shows the current watchlist, daily close semantics, API/Mongo health, source freshness, and the areas where persisted analytics appear.")
    add_screenshot(doc, "doc-entity-deep-dive.png", "Entity Deep Dive — Apple",
                   "The Entity workspace demonstrates selection by configured entity, price context, change versus the previous close, data mode, and provisional-bar labeling.")
    add_screenshot(doc, "doc-live-firehose.png", "Live Firehose — WebSocket status",
                   "The Firehose workspace shows an open WebSocket, heartbeat timing, filter controls, and the distinction between server-replayed snapshots and new arrivals.")
    add_screenshot(doc, "doc-pipeline.png", "Pipeline & Methodology — operational transparency",
                   "The Pipeline workspace exposes recorded operational status and deliberately labels uninstrumented fields as UNKNOWN instead of inventing health.")

    doc.add_heading("8. Operations and troubleshooting", level=1)
    table(doc, ["Symptom", "First checks", "Interpretation or action"], [
        ["API unreachable", "docker compose ps; curl.exe http://localhost:8000/health", "Check API logs and port 8000."],
        ["Data stale", "GET /api/v1/sources/status; ingestion logs", "Inspect provider throttling, feed status, or network access."],
        ["News source failed", "Source status and ingestion-news logs", "One source can fail without stopping other feeds; verify URL and retry behavior."],
        ["Empty analytics", "analytics-job logs and interval setting", "Wait for a cycle; anomalies also need historical windows and volume."],
        ["Firehose no events", "WebSocket status and spark-text logs", "A connected socket can legitimately have no newly matched event."],
        ["Fresh build has no data", "Wait for init jobs, then run smoke scripts", "Healthy containers alone do not prove end-to-end data flow."],
    ])
    doc.add_heading("9. Verification snapshot", level=1)
    table(doc, ["Check", "Observed result"], [
        ["Dashboard", "HTTP 200 and rendered MarketPulse workspaces."],
        ["API and Mongo", "API ok; Mongo ok."],
        ["Market persistence", "2,746 market bars across the configured watchlist."],
        ["News persistence", "947 news articles after Google News entity feeds were enabled."],
        ["Social persistence", "71 Bluesky posts at the documentation capture time."],
        ["Enabled source health", "yfinance, Bluesky, CNBC, MarketWatch, Mint, Business Standard, and Google News successful."],
        ["Retired source handling", "Yahoo Finance RSS disabled and surfaced as DISABLED, not as an active failure."],
    ])

    doc.add_heading("10. Limitations and responsible interpretation", level=1)
    for item in [
        "Daily market bars are not live ticks and may be provisional for the current trading date.",
        "External feeds can change, rate-limit, redirect, or become unavailable.",
        "Bluesky is a single public platform and is not a representative sample of investors.",
        "VADER sentiment is a lightweight heuristic; news scoring uses headlines.",
        "At-least-once delivery plus idempotent writes is not exactly-once processing.",
        "The local stack has no authentication and should not be exposed publicly without hardening.",
        "All analytics are descriptive and must not be presented as investment advice or causal evidence.",
    ]:
        bullet(doc, item)

    doc.add_heading("Closing demo statement", level=1)
    doc.add_paragraph(
        "MarketPulse demonstrates a complete, observable data product rather than only a dashboard mock-up: real providers feed Kafka, "
        "Spark validates and persists records, analytics enrich the data, FastAPI exposes typed read models, and React communicates freshness "
        "and uncertainty directly to the user. Each team member can explain one layer while understanding how the layers connect."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("MarketPulse Project Handoff | Academic project | Descriptive analytics only").font.size = Pt(8)
    doc.core_properties.title = "MarketPulse Project Handoff and Demo Guide"
    doc.core_properties.subject = "Complete architecture, implementation, ownership split, and dashboard evidence"
    doc.core_properties.author = "MarketPulse team"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
